import re
from pathlib import Path
from typing import Tuple
from fastapi import HTTPException
import pymupdf
import docx

def extract_text_from_pdf(file_path: Path) -> str:
    """Extracts plain text from a PDF file using PyMuPDF."""
    text_content = []
    try:
        doc = pymupdf.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            if text:
                text_content.append(text)
        doc.close()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PDF file: {str(e)}")
        
    full_text = "\n".join(text_content).strip()
    if not full_text or len(full_text.replace(" ", "").replace("\n", "")) < 20:
        raise HTTPException(
            status_code=422,
            detail="The uploaded PDF contains no extractable text (it may be a scanned image). "
                   "OCR is not supported in this version. Please upload a searchable PDF, DOCX, or TXT file."
        )
    return full_text

def extract_text_from_docx(file_path: Path) -> str:
    """Extracts plain text from a DOCX file using python-docx."""
    try:
        doc = docx.Document(file_path)
        text_content = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text_content.append(row_text)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read DOCX file: {str(e)}")
        
    full_text = "\n".join(text_content).strip()
    if not full_text:
        raise HTTPException(status_code=422, detail="The uploaded DOCX file is empty.")
    return full_text

def extract_text_from_txt(file_path: Path) -> str:
    """Extracts plain text from a TXT file."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read().strip()
    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="latin-1") as f:
                text = f.read().strip()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to decode text file: {str(e)}")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read TXT file: {str(e)}")
        
    if not text:
        raise HTTPException(status_code=422, detail="The uploaded TXT file is empty.")
    return text

def parse_resume_file(file_path: Path) -> Tuple[str, str]:
    """
    Determines file type and extracts normalized text and a short preview.
    """
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(file_path)
    elif ext == ".txt":
        raw_text = extract_text_from_txt(file_path)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file format: {ext}")
        
    # Clean excessive whitespace
    normalized_text = re.sub(r"\n{3,}", "\n\n", raw_text)
    preview = normalized_text[:500] + ("..." if len(normalized_text) > 500 else "")
    return normalized_text, preview
